"""Генерация отчётов: «Отчет о проведенных заседаниях СНО» (.docx) и выгрузка участий (.xlsx).

DOCX строится по шаблону templates/sno_meetings_template.docx (очищенная копия
исходного документа): строка-заголовок таблицы + ОДНА строка-прототип с
плейсхолдерами {N} {DATE} {TIME} {LOCATION} {TOPIC} {FORMAT}. Для каждого заседания
строка-прототип клонируется (deepcopy XML), поэтому шрифты, границы, ширины колонок
и выравнивание в точности как в оригинале. Абзацы заголовка, «Приложение …» и
подписей заполняются заменой текста внутри существующих runs (форматирование
сохраняется); абзац подписи клонируется для каждого подписанта.
"""

from __future__ import annotations

import copy
import io
from datetime import date
from pathlib import Path
from typing import Any, Iterable, Optional

TEMPLATE_PATH = Path(__file__).resolve().parent / "templates" / "sno_meetings_template.docx"
SIGNATURE_LINE = "_______________"


def _to_date(value: Any) -> date:
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value)[:10])


def _replace_in_runs(paragraph, mapping: dict[str, str]) -> None:  # noqa: ANN001
    """Replace placeholders run-by-run (each placeholder lives inside one run)."""
    for run in paragraph.runs:
        t = run.text
        new = t
        for k, v in mapping.items():
            if k in new:
                new = new.replace(k, v)
        if new != t:
            run.text = new  # '\n' → <w:br/>, formatting (rPr) untouched


def _set_paragraph_text(paragraph, value: str) -> None:  # noqa: ANN001
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(value)
        return
    runs[0].text = value
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def _format_lines(fmt: str) -> str:
    """'Форсайт-сессия, Собрание' → 'Форсайт-сессия,\nСобрание' (как в образце)."""
    items = [s.strip() for s in (fmt or "").split(",") if s.strip()]
    return ",\n".join(items)


def build_meetings_docx(
    meetings: Iterable[dict],
    year: int,
    sno_name: str,
    appendix_label: str = "Приложение Е",
    signatories: Optional[list[dict]] = None,
    template_path: Optional[Path | str] = None,
) -> bytes:
    """Return .docx bytes. meetings: dicts with meeting_date, meeting_time, location, topic, format
    (already sorted; numbering 1..N is assigned here in the given order)."""
    from docx import Document

    doc = Document(str(template_path or TEMPLATE_PATH))

    # Заголовок и «Приложение»
    for p in doc.paragraphs:
        if "{APPENDIX}" in p.text:
            _replace_in_runs(p, {"{APPENDIX}": appendix_label or ""})
        if "{SNO_NAME}" in p.text or "{YEAR}" in p.text:
            _replace_in_runs(p, {"{SNO_NAME}": sno_name or "", "{YEAR}": str(year)})

    # Таблица: клонируем строку-прототип
    table = doc.tables[0]
    tbl = table._tbl
    proto = tbl.tr_lst[1]
    anchor = proto
    for i, m in enumerate(meetings, start=1):
        tr = copy.deepcopy(proto)
        anchor.addnext(tr)
        anchor = tr
        from docx.table import _Row

        row = _Row(tr, table)
        cells = row.cells
        d = _to_date(m["meeting_date"]).strftime("%d.%m.%Y")
        t = (m.get("meeting_time") or "").strip()
        _set_paragraph_text(cells[0].paragraphs[0], f"{i}.")
        date_p, time_p = cells[1].paragraphs[0], cells[1].paragraphs[1]
        _set_paragraph_text(date_p, f"{d}," if t else d)
        if t:
            _set_paragraph_text(time_p, t)
        else:
            time_p._p.getparent().remove(time_p._p)
        _set_paragraph_text(cells[2].paragraphs[0], (m.get("location") or "").strip())
        _set_paragraph_text(cells[3].paragraphs[0], (m.get("topic") or "").strip())
        _set_paragraph_text(cells[4].paragraphs[0], _format_lines(m.get("format") or ""))
    tbl.remove(proto)

    # Подписи: клонируем абзац-прототип для каждого подписанта
    sig_proto = next(p for p in doc.paragraphs if "{POSITION}" in p.text)
    sigs = [s for s in (signatories or []) if (s.get("position") or s.get("name"))]
    anchor = sig_proto._p
    for idx, s in enumerate(sigs):
        if idx > 0:
            # два пустых абзаца с тем же форматированием, как в образце
            for _ in range(2):
                empty = copy.deepcopy(sig_proto._p)
                for r in empty.findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r"):
                    empty.remove(r)
                anchor.addnext(empty)
                anchor = empty
        sp = copy.deepcopy(sig_proto._p)
        anchor.addnext(sp)
        anchor = sp
        from docx.text.paragraph import Paragraph

        para = Paragraph(sp, sig_proto._parent)
        pos = (s.get("position") or "").strip()
        _replace_in_runs(para, {
            "{POSITION} ": f"{pos} " if pos else "",
            "{POSITION}": pos,
            "{NAME}": (s.get("name") or "").strip(),
        })
    sig_proto._p.getparent().remove(sig_proto._p)

    doc.core_properties.title = f"Отчет о проведенных заседаниях СНО «{sno_name}» в {year} году"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_year_xlsx(participations: list[dict], meetings: Optional[list[dict]] = None) -> bytes:
    """Excel: лист «Участия» (+ лист «Заседания», если переданы)."""
    import pandas as pd

    part_df = pd.DataFrame(
        [
            {
                "ФИО": r["full_name"],
                "Логин": r["login"],
                "Мероприятие": r["title"],
                "Тип": r["type"],
                "Дата": str(r["event_date"]),
                "Номер в портфолио": r.get("achievement_number") or "",
                "Есть номер": "да" if r.get("achievement_number") else "нет",
                "Тема статьи": r.get("article_topic") or "",
                "Индексация": r.get("indexing") or "",
                "Добавлено": str(r["joined_at"])[:19],
            }
            for r in participations
        ],
        columns=["ФИО", "Логин", "Мероприятие", "Тип", "Дата", "Номер в портфолио",
                 "Есть номер", "Тема статьи", "Индексация", "Добавлено"],
    )
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        part_df.to_excel(xw, sheet_name="Участия", index=False)
        _autowidth(xw.sheets["Участия"])
        if meetings is not None:
            m_df = pd.DataFrame(
                [
                    {
                        "№": i,
                        "Вид": m.get("kind") or "",
                        "Дата": _to_date(m["meeting_date"]).strftime("%d.%m.%Y"),
                        "Время": m.get("meeting_time") or "",
                        "Локация": m.get("location") or "",
                        "Тема": m.get("topic") or "",
                        "Формат": m.get("format") or "",
                    }
                    for i, m in enumerate(meetings, start=1)
                ],
                columns=["№", "Вид", "Дата", "Время", "Локация", "Тема", "Формат"],
            )
            m_df.to_excel(xw, sheet_name="Заседания", index=False)
            _autowidth(xw.sheets["Заседания"])
    return buf.getvalue()


def _autowidth(ws) -> None:  # noqa: ANN001
    for col in ws.columns:
        width = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 60)
