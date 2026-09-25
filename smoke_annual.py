"""Smoke tests for the annual report (catalog, achievements, counting, migration, .docx).
Run through smoke_test.py (SQLite by default, Postgres via SMOKE_DATABASE_URL)."""

from __future__ import annotations

import hashlib
import io
import json

from sqlalchemy import text

import achievements as ach
import annual_fixtures as fx
import annual_report
import db
import report

NEW_TABLES = ("achievement_people", "achievements", "kind_subpoints", "kinds", "indicator_rows", "indicators")
ROWS_PER_INDICATOR = [4, 5, 4, 4, 4, 4, 4, 5, 4, 2, 4, 4, 8, 4, 4, 1, 0, 0]


def _admin(target) -> dict:  # noqa: ANN001
    return next(u for u in db.list_users(db_path=target) if u["role"] == "admin")


def run_catalog(target) -> None:  # noqa: ANN001
    cat = ach.catalog(target)
    assert [i["label"] for i in cat] == [x[1] for x in ach.INDICATORS]
    assert [len(i["rows"]) for i in cat] == ROWS_PER_INDICATOR
    assert [r["label"] for r in cat[1]["rows"]][4] == "- статья в журнале из перечня Белый список"
    with db.get_engine(target).begin() as conn:
        assert ach.seed_catalog(conn) == 0  # idempotent
    db.init_db(db_path=target)
    with db.get_engine(target).connect() as conn:
        assert conn.execute(text("SELECT COUNT(*) FROM indicators")).scalar_one() == 18
        assert conn.execute(text("SELECT COUNT(*) FROM indicator_rows")).scalar_one() == sum(ROWS_PER_INDICATOR)
        assert conn.execute(text("SELECT COUNT(*) FROM kinds")).scalar_one() == len(ach.KINDS)
    # admin edits survive restarts (seed never overwrites)
    row = cat[13]["rows"][3]  # 14 / городской
    ach.update_row(row["id"], "- городской уровень (г. Ростов-на-Дону)", True, db_path=target)
    ach.move_indicator(cat[17]["id"], -1, db_path=target)
    db.init_db(db_path=target)
    cat2 = ach.catalog(target)
    r2 = next(r for r in cat2[13]["rows"] if r["id"] == row["id"])
    assert r2["label"].endswith("(г. Ростов-на-Дону)") and r2["hidden"] == 1
    assert cat2[16]["code"] == "I18" and cat2[17]["code"] == "I17"
    ach.move_indicator(cat[17]["id"], +1, db_path=target)
    ach.update_row(row["id"], row["label"], False, db_path=target)
    new_id = ach.add_indicator("19. Тестовый показатель", "level", db_path=target)
    rid = ach.add_row(new_id, "- международный уровень", db_path=target)
    assert any(k["indicator_id"] == new_id for k in ach.list_kinds(target))
    ach.delete_row(rid, db_path=target)
    ach.delete_indicator(new_id, db_path=target)
    assert len(ach.catalog(target)) == 18
    # never hard-delete with records
    admin = _admin(target)
    m = db.create_user("cat_m", "pass12345", "Каталогов К.К.", db_path=target)
    K = ach.kind_by_code("contest", target)
    aid = ach.save_achievement({"kind_id": K["id"], "owner_id": m, "row_id": cat[2]["rows"][0]["id"],
                                "title": "Конкурс X", "date_from": "2025-05-05", "result": "участие"},
                               admin, db_path=target)["id"]
    for fn, arg in ((ach.delete_row, cat[2]["rows"][0]["id"]), (ach.delete_indicator, cat[2]["id"])):
        try:
            fn(arg, db_path=target)
            raise AssertionError("delete with records must fail")
        except ValueError:
            pass
    ach.delete_achievement(aid, db_path=target)
    db.delete_user(m, db_path=target)
    print("  annual catalog OK (18 indicators verbatim, idempotent seed, admin edits kept, no delete with records)")


def run_numbers_and_validation(target) -> None:  # noqa: ANN001
    assert ach.normalize_number(" № P-H-1700-25 ") == "Р-Н-1700-25"
    assert ach.normalize_number("Nº р-х-349-25") == "Р-Х-349-25"
    assert ach.normalize_number("Р – О – 805 – 25") == "Р-О-805-25"
    assert ach.normalize_number("") is None
    assert ach.academic_year_range("2025-2026") == ("2025-09-01", "2026-08-31")
    for bad in ("2025", "2025-2027"):
        try:
            ach.academic_year_range(bad)
            raise AssertionError(bad)
        except ValueError:
            pass
    ok = [{"name": "Б", "share": 48}, {"name": "В", "share": 1}]
    assert ach.validate_shares(49, ok) == 98
    for owner, co, msg in ((None, [], "главного"), (0, [], "больше 0"), (101, [], "больше 0"),
                           (50, [{"name": "Б", "share": None}], "укажите долю"),
                           (50, [{"name": "", "share": 10}], "ФИО"),
                           (60, [{"name": "Б", "share": 50}], "больше 100")):
        try:
            ach.validate_shares(owner, co)
            raise AssertionError((owner, co))
        except ValueError as e:
            assert msg in str(e), (msg, str(e))
    print("  numbers normalized (Latin P-H → Р-Н, №/Nº stripped), shares validated")


def run_fixtures_report(target) -> None:  # noqa: ANN001
    res = fx.load(target)
    uid = res["users"]
    data = ach.report_data(2025, db_path=target)
    counts = {r["label"] and (ind["code"], i): r["count"]
              for ind in data["indicators"] for i, r in enumerate(ind["rows"])}
    exp = {("I01", 0): 5, ("I01", 1): 0, ("I01", 2): 1, ("I02", 0): 1, ("I02", 2): 1, ("I02", 3): 1,
           ("I02", 4): 1, ("I03", 1): 3, ("I03", 3): 2, ("I04", 0): 1, ("I04", 1): 2, ("I07", 1): 1,
           ("I08", 0): 1, ("I08", 1): 3, ("I08", 2): 1, ("I08", 4): 1, ("I09", 0): 2, ("I09", 1): 1,
           ("I09", 3): 1, ("I10", 0): 2, ("I11", 0): 2, ("I12", 2): 1, ("I15", 1): 2, ("I16", 0): 1}
    for k, v in counts.items():
        assert v == exp.get(k, 0), (k, v, exp.get(k, 0))
    totals = {i["code"]: i["count"] for i in data["indicators"]}
    assert totals["I01"] == 6 and totals["I02"] == 4 and totals["I17"] == 1 and totals["I18"] == 1
    assert totals["I05"] == 0 and totals["I13"] == 0
    assert data["total"] == 39 and data["zaochno"] == 1 and data["warnings"] == []
    # several dokladov by one person at one event: one event line, two numbered lines
    items = next(i for i in data["indicators"] if i["code"] == "I01")["rows"][0]["items"]
    ev = [x for x in items if x[0] == "event"]
    assert len(ev) == 2 and sum(1 for x in items if x[0] == "name" and x[1] == fx.PEOPLE["starostin"]) == 3
    forum_i = next(n for n, x in enumerate(items) if x[0] == "event" and x[1].startswith("Международный"))
    forum_lines = [x for x in items[forum_i + 1:] if x[0] == "name"]
    assert len(forum_lines) == 3 and all("доклад на тему: «" in x[2] for x in forum_lines)
    assert any(x[2].endswith("Р-Н-5638-25") for x in forum_lines)  # Latin P-H normalized
    # publication line: shares in author order, co-authors printed, counted once
    pub = next(i for i in data["indicators"] if i["code"] == "I02")
    vak = pub["rows"][2]["items"][0][1]
    assert "Доля участия: 49%, 48%, 1%." in vak and "Шевченко Виктория Николаевна" in vak and "Р-Н-6100-25" in vak
    rinc = pub["rows"][3]["items"][0][1]
    assert rinc.startswith("Мартынюк И.О., Старостин Д.В. Методы") and "DOI: 10.23947/interagro.2025.117-119." in rinc
    assert "Доля участия: 50%, 50%." in rinc
    scopus = pub["rows"][0]["items"][0][1]
    assert scopus.startswith("Starostin, D.") and "Доля участия: 10%, 10%, 10%." in scopus
    # per-person: publication counts only for the main author
    per = {r["login"]: r for r in ach.stats_by_person(2025, db_path=target)}
    assert per["martynuk"]["by_kind"].get("Публикация") == 1
    assert "Публикация" not in per["katanaeva"]["by_kind"]
    # contest result printed at the end of the line; «участие» is not printed
    univ = next(i for i in data["indicators"] if i["code"] == "I03")["rows"][3]["items"]
    lines3 = [x[1] for x in univ]
    assert any(x.endswith("Р-П-1444-25, диплом II степени") for x in lines3), lines3
    assert not any("участие" in x for x in lines3), lines3
    # publication «Без индексации»: valid record, not printed, no warning
    admin_u = next(u for u in db.list_users(db_path=target) if u["role"] == "admin")
    pub_k = next(k for k in ach.list_kinds(target, admin=True) if k["form"] == "publication")
    nid = ach.save_achievement({"kind_id": pub_k["id"], "owner_id": uid["gaidai"], "row_id": ach.NO_INDEX,
                                "title": "Статья в сборнике без индексации", "journal": "Сборник",
                                "year": 2025, "owner_share": 100.0}, admin_u, db_path=target)["id"]
    rec = ach.get_achievement(nid, db_path=target)
    assert rec["details"].get("no_index") and rec["row_id"] is None and not rec["counted"] and rec["issues"] == []
    data2 = ach.report_data(2025, db_path=target)
    assert data2["total"] == 39 and data2["warnings"] == []
    ach.delete_achievement(nid, db_path=target)
    assert per["starostin"]["by_kind"]["Публикация"] == 1  # own Scopus article only
    # year filter: stipend 2024-2025 in 2024 and 2025; 2025-2026 in 2025 and 2026
    s24 = [r for r in ach.list_achievements(year=2024, db_path=target) if r["form"] == "stipend"]
    s26 = [r for r in ach.list_achievements(year=2026, db_path=target) if r["form"] == "stipend"]
    assert len(s24) == 2 and len(s26) == 4
    # duplicate publication (title+year / DOI) → refused with the owner's name
    K = ach.kind_by_code("publication", target)
    member = db.get_user_by_id(uid["katanaeva"], db_path=target)
    base = {"kind_id": K["id"], "row_id": ach.row_by_code("I02.rinc", target)["id"], "owner_share": 100,
            "journal": "Ж", "year": 2025}
    for dup in ({"title": "  методы внесения ПРОБИОТИКОВ в комбикорма: обзор "},
                {"title": "Другое название", "doi": "https://doi.org/10.23947/INTERAGRO.2025.117-119"}):
        try:
            ach.save_achievement({**base, **dup}, member, db_path=target)
            raise AssertionError("duplicate must be refused")
        except ach.DuplicateAchievement as e:
            assert fx.PEOPLE["martynuk"] in str(e)
    try:
        ach.save_achievement({**base, "title": "Новая", "owner_share": 70,
                              "coauthors": [{"name": "Внешний А.А.", "share": 40}]}, member, db_path=target)
        raise AssertionError("total > 100 must be refused")
    except ValueError as e:
        assert "больше 100" in str(e)
    # member rights: no admin kinds, only own records
    try:
        ach.save_achievement({"kind_id": ach.kind_by_code("agreement", target)["id"], "title": "x",
                              "agreement": "1", "date_from": "2025-01-01"}, member, db_path=target)
        raise AssertionError("member must not add admin kinds")
    except ValueError:
        pass
    other = ach.list_achievements(owner_id=uid["martynuk"], db_path=target)[0]
    try:
        ach.save_achievement({"kind_id": other["kind_id"], "title": "x"}, member, achievement_id=other["id"],
                             db_path=target)
        raise AssertionError("member must not edit others")
    except ValueError:
        pass
    assert not ach.delete_achievement(other["id"], owner_id=member["id"], db_path=target)
    # member stats see no names
    ov = ach.events_overview(2025, db_path=target)
    assert ov and all(set(e) == {"title", "kind", "date", "records", "people"} for e in ov)
    blob = json.dumps(ov, ensure_ascii=False)
    assert not any(n in blob for n in fx.PEOPLE.values())

    # .docx: landscape, title, every indicator/sub-row verbatim with counts, conclusion, signatures
    b = annual_report.build_annual_docx(data, ach.get_conclusion(2025, db_path=target),
                                        db.get_report_settings(db_path=target)["signatories"])
    check_docx(b, data)  # conclusion/signatures from fixtures
    x = annual_report.build_achievements_xlsx(data, db.list_meetings(2025, db_path=target))
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(x))
    assert wb.sheetnames == ["Достижения", "Показатели", "Заседания"]
    assert wb["Достижения"].max_row == len(data["records"]) + 1
    # the existing meetings report («Приложение Е») still generates alongside
    mid = db.add_meeting("2025-03-01", "13:30", "Ауд. 1", "Заседание СНО", "Собрание", db_path=target)
    mb = report.build_meetings_docx(db.list_meetings(2025, db_path=target), 2025, fx.SNO_NAME,
                                    db.DEFAULT_APPENDIX_LABEL, fx.SIGNATORIES)
    from docx import Document

    mdoc = Document(io.BytesIO(mb))
    mtext = "\n".join(p.text for p in mdoc.paragraphs)
    assert "Приложение Е" in mtext and "Сельское хозяйство" in mtext
    assert any("Заседание СНО" in c.text for t in mdoc.tables for row in t.rows for c in row.cells)
    db.delete_meeting(mid, db_path=target)
    print(f"  annual report OK: {data['total']} units, docx {len(b)} bytes, xlsx {len(x)} bytes; "
          "meetings report still OK")


def check_docx(b: bytes, data: dict) -> None:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.oxml.ns import qn

    d = Document(io.BytesIO(b))
    sec = d.sections[0]
    assert sec.orientation == WD_ORIENT.LANDSCAPE and sec.page_width > sec.page_height
    paras = [p.text for p in d.paragraphs]
    assert paras[0] == f"Отчет о работе СНО «{data['sno_name']}»" and paras[1] == f"За {data['year']} год"
    t0, t1 = d.tables
    assert [c.text for c in t0.rows[0].cells] == ["Показатель", "Количество",
                                                  "Отметка о выполнении с указанием детальной информации"]
    rows = [r for t in (t0, t1) for r in t.rows][1:]
    expected = []
    for ind in data["indicators"]:
        expected.append((ind["label"], str(ind["count"]), True))
        expected += [(r["label"], str(r["count"]) if r["count"] else "", False) for r in ind["rows"]]
    assert len(rows) == len(expected), (len(rows), len(expected))
    for r, (label, cnt, bold) in zip(rows, expected):
        cells = r.cells
        assert cells[0].text == label, (cells[0].text, label)
        assert cells[1].text == cnt, (label, cells[1].text, cnt)
        assert bool(cells[0].paragraphs[0].runs[0].bold) == bold
    assert len(t0.rows) == 1 + sum(1 + len(i["rows"]) for i in data["indicators"][:5])
    # numbered detail lines restart per sub-row (separate w:num per cell)
    nums = []
    for r in rows:
        ids = {p._p.find(qn("w:pPr") + "/" + qn("w:numPr") + "/" + qn("w:numId")).get(qn("w:val"))
               for p in r.cells[2].paragraphs
               if p._p.find(qn("w:pPr") + "/" + qn("w:numPr")) is not None}
        assert len(ids) <= 1
        nums += list(ids)
    assert len(nums) == len(set(nums)) and nums
    body = "\n".join(paras)
    assert fx.CONCLUSION in body
    for sig in fx.SIGNATORIES:
        assert f"{sig['position']} {annual_report.SIGNATURE_LINE}/ {sig['name']}" in body
    assert [c.width for c in t0.rows[0].cells] == [c.width for c in t1.rows[0].cells]  # template grid


def run_orphans_and_meetings(target) -> None:  # noqa: ANN001
    admin = _admin(target)
    u = db.create_user("orph_a", "pass12345", "Сиротин С.С.", db_path=target)
    K = ach.kind_by_code("doklad", target)
    r = ach.row_by_code("I01.reg", target)["id"]
    a1 = ach.save_achievement({"kind_id": K["id"], "owner_id": u, "row_id": r, "title": "Форум Сирот",
                               "date_from": "2025-09-09", "topic": "Т1", "ochno": True}, admin, db_path=target)["id"]
    a2 = ach.save_achievement({"kind_id": K["id"], "owner_id": u, "row_id": r, "title": "форум сирот ",
                               "date_from": "2025-09-09", "topic": "Т2", "ochno": True}, admin, db_path=target)["id"]
    e1 = ach.get_achievement(a1, target)["event_id"]
    assert e1 and e1 == ach.get_achievement(a2, target)["event_id"]  # same event, 2 dokladov
    offered = {e["event_id"]: e for e in db.unlinked_member_events(2025, db_path=target)}
    assert offered[e1]["participants_count"] == 1 and offered[e1]["suggested_kind"] == "Конференция"
    mid = db.add_meeting("2025-09-09", "13:30", "Ауд", "Выступления", "Конференция", kind="Конференция",
                         event_id=e1, db_path=target)
    assert ach.delete_achievement(a1, db_path=target)
    assert db.get_meeting(mid, db_path=target)["event_id"] == e1  # still used by a2
    assert ach.delete_achievement(a2, db_path=target)
    with db.get_engine(target).connect() as conn:
        assert conn.execute(text("SELECT COUNT(*) FROM events WHERE id = :e"), {"e": e1}).scalar_one() == 0
    m = db.get_meeting(mid, db_path=target)
    assert m is not None and m["event_id"] is None  # meeting row kept in the report
    db.delete_meeting(mid, db_path=target)
    db.delete_user(u, db_path=target)
    print("  achievements: event grouping, orphan cleanup, meeting row kept OK")


def _md5(target, tables=("users", "events", "participations", "meetings", "settings")) -> dict:  # noqa: ANN001
    out = {}
    with db.get_engine(target).connect() as conn:
        for t in tables:
            rows = conn.execute(text(f"SELECT * FROM {t} ORDER BY 1")).all()
            out[t] = hashlib.md5(repr([tuple(str(v) for v in r) for r in rows]).encode()).hexdigest()
    return out


def run_migration(target) -> None:  # noqa: ANN001
    """DB of the current release (no achievement tables) with data → init_db migrates it."""
    import auth

    db.init_db(db_path=target, seed_admin=True)
    with db.get_engine(target).begin() as conn:
        for t in NEW_TABLES:
            conn.execute(text(f"DROP TABLE {t}"))
    a = db.create_user("mig_a", "pass12345", "Первый Автор", db_path=target)
    b = db.create_user("mig_b", "pass12345", "Второй Автор", db_path=target)
    c = db.create_user("mig_c", "pass12345", "Третий Член", db_path=target)
    add = lambda u, *a_, **k: db.add_participation_ex(u, *a_, db_path=target, **k)  # noqa: E731
    add(a, "Конференция ИТНО", "конференция", "2025-09-10", achievement_number="P-H-1-25")
    add(b, "Конференция ИТНО", "конференция", "2025-09-10", achievement_number="Р-Н-2-25")
    add(a, "Статья про корма", "статья", "2025-05-01", article_topic="корма", indexing="ВАК")
    add(b, "Статья про корма", "статья", "2025-05-01", article_topic="корма", indexing="ВАК",
        achievement_number="Р-Н-3-25")
    add(c, "Статья без индекса", "статья", "2025-06-01", indexing="Без индексации")
    add(b, "Кейс-чемпионат", "конкурс", "2025-04-04", achievement_number="Р-Н-4-25")
    add(c, "Стипендия Губернатора", "стипендия", "2025-02-01")
    add(a, "УМНИК 2025", "грант", "2025-03-03")
    eid = db.get_or_create_event("Конференция ИТНО", "конференция", "2025-09-10", db_path=target)
    db.add_meeting("2025-09-10", "13:30", "Ауд", "ИТНО", "Конференция", kind="Конференция", event_id=eid,
                   db_path=target)
    db.save_report_settings("Сельское хозяйство", "Приложение Е", fx.SIGNATORIES, db_path=target)
    tok = auth.make_auth_token(db.get_user_by_id(a, db_path=target), db_path=target)
    before = _md5(target)
    with db.get_engine(target).connect() as conn:
        n_parts = conn.execute(text("SELECT COUNT(*) FROM participations")).scalar_one()

    db.init_db(db_path=target)  # migration
    db.init_db(db_path=target)  # idempotent
    assert _md5(target) == before, "legacy tables must stay byte-identical"
    recs = ach.list_achievements(db_path=target)
    with db.get_engine(target).connect() as conn:
        n_people = conn.execute(text("SELECT COUNT(*) FROM achievement_people")).scalar_one()
    assert len(recs) + n_people == n_parts == 8, (len(recs), n_people, n_parts)
    by = {(r["kind_code"], r["owner_login"], r["title"]): r for r in recs}
    d1 = by[("doklad", "mig_a", "Конференция ИТНО")]
    assert d1["number"] == "Р-Н-1-25" and d1["ochno"] == 1 and d1["row_id"] is None and d1["event_id"] == eid
    assert "укажите уровень" in d1["issues"] and not d1["counted"]
    art = by[("publication", "mig_a", "Статья про корма")]  # main author = first added
    assert art["row_code"] == "I02.vak" and [p["name"] for p in art["people"]] == ["Второй Автор"]
    assert art["details"]["coauthor_numbers"] == ["Р-Н-3-25"] and "заполните долю" in art["issues"]
    assert art["counted"]
    noidx = by[("publication", "mig_c", "Статья без индекса")]
    assert noidx["row_id"] is None and noidx["details"].get("no_index") and not noidx["counted"]
    assert not any("индексац" in i for i in noidx["issues"])
    assert by[("contest", "mig_b", "Кейс-чемпионат")]["number"] == "Р-Н-4-25"
    assert by[("stipend", "mig_c", "Стипендия Губернатора")]["row_code"] == "I08.other"
    g = by[("grant", "mig_a", "УМНИК 2025")]
    assert g["subpoint_id"] is None and "выберите подпункт гранта" in g["issues"]
    data = ach.report_data(2025, db_path=target)
    assert len(data["warnings"]) == 6 and data["total"] == 2  # article (ВАК) + stipend «иные»
    # users, roles, cookies, meetings keep working
    assert auth.user_from_token(tok, db_path=target)["id"] == a
    assert db.list_meetings(2025, db_path=target)[0]["event_id"] == eid
    assert db.get_user_by_login("mig_c", db_path=target)["role"] == "member"
    assert auth.verify_password("pass12345", db.get_user_by_login("mig_b", db_path=target)["password_hash"])
    # owner fills the shares → warning gone; deleting a migrated record removes its legacy rows
    admin = _admin(target)
    ach.save_achievement({"kind_id": art["kind_id"], "owner_id": a, "row_id": art["row_id"],
                          "title": art["title"], "journal": "Вестник", "year": 2025, "owner_share": 60,
                          "coauthors": [{"user_id": b, "name": "Второй Автор", "share": 40}]},
                         admin, achievement_id=art["id"], db_path=target)
    art2 = ach.get_achievement(art["id"], target)
    assert art2["issues"] == [] and ach.shares_text(art2) == "доля участия: 60%, 40%"
    db.init_db(db_path=target)
    assert len(ach.list_achievements(db_path=target)) == len(recs)  # nothing re-imported
    assert ach.delete_achievement(d1["id"], db_path=target)
    db.init_db(db_path=target)
    assert len(ach.list_achievements(db_path=target)) == len(recs) - 1
    assert db.list_meetings(2025, db_path=target)[0]["event_id"] == eid  # b's doklad still uses it
    print(f"  migration from current schema OK ({n_parts} participations → {len(recs)} records "
          f"+ {n_people} co-author; legacy md5 unchanged; idempotent; warnings={len(data['warnings'])})")


def run_all(target, fresh) -> None:  # noqa: ANN001
    """target: DB for catalog/fixture tests (fresh); fresh(): returns a new empty DB target."""
    db.init_db(db_path=target, seed_admin=True)
    run_catalog(target)
    run_numbers_and_validation(target)
    run_orphans_and_meetings(target)
    run_fixtures_report(target)
    run_migration(fresh())
